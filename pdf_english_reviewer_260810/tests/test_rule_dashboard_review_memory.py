import unittest
import os
import tempfile
from pathlib import Path

os.environ.setdefault(
    "REVIEWER_DB_PATH",
    str(Path(tempfile.gettempdir()) / f"pdf_english_reviewer_dashboard_{os.getpid()}.db"),
)

from app.exports.rule_excel import build_rule_dashboard_workbook, safe_excel_value
from app.review_memory.service import parse_memory_txt, safe_backup_path, write_memory_txt
import app.main as main


class RuleDashboardReviewMemoryTests(unittest.TestCase):
    def test_requested_routes_are_registered(self):
        routes = {route.path for route in main.app.routes}
        for path in (
            "/api/glossary/export.xlsx",
            "/api/glossary/export.txt",
            "/api/glossary/export.db",
            "/api/engines/status",
            "/api/engines/vale/status",
            "/api/engines/ollama/status",
        ):
            self.assertIn(path, routes)

    def test_excel_formula_injection_values_are_escaped(self):
        for value in ("=SUM(A1:A2)", "+cmd", "-cmd", "@name"):
            self.assertTrue(str(safe_excel_value(value)).startswith("'"))
        self.assertEqual(safe_excel_value("plain text"), "plain text")

    def test_rule_dashboard_xlsx_contains_required_sheets(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "rule_dashboard.xlsx"
            build_rule_dashboard_workbook(
                path,
                summary={"total_rules": 1, "total_findings": 1},
                inventory=[{"rule_key": "basic:BuiltIn:R1", "rule_id": "R1", "rule_name": "Rule 1"}],
                findings=[{"document": "manual.pdf", "rule_key": "basic:BuiltIn:R1", "message": "=safe"}],
                statistics=[{"rule_key": "basic:BuiltIn:R1", "trigger_count": 1}],
                criteria={"export_type": "current"},
            )
            from openpyxl import load_workbook
            workbook = load_workbook(path, read_only=True)
            try:
                self.assertEqual(
                    set(workbook.sheetnames),
                    {"Summary", "Rule Inventory", "Rule Findings", "Rule Statistics", "Export Criteria"},
                )
            finally:
                workbook.close()

    def test_review_memory_txt_round_trips_escaped_values(self):
        memory = {
            "manifest": {"schema_version": 1, "created_at": "2026-07-13T00:00:00Z"},
            "rules": [{"rule_id": "R1", "pattern": "line\nwith\ttab\\slash"}],
            "glossary": [],
            "exceptions": [],
            "feedback": [],
            "presets": [],
            "statistics": [],
        }
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "review_memory.txt"
            write_memory_txt(path, memory)
            sections = parse_memory_txt(path.read_text(encoding="utf-8"))
            self.assertEqual(sections["RULES"][0]["pattern"], "line\nwith\ttab\\slash")

    def test_review_memory_backup_id_rejects_path_traversal(self):
        with self.assertRaises(ValueError):
            safe_backup_path(Path("data/review_memory"), "../bad")


if __name__ == "__main__":
    unittest.main()
