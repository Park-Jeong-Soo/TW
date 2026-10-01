from __future__ import annotations

import hashlib
import json
import re
import shutil
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from app.exports.rule_excel import build_memory_workbook, safe_excel_value


SCHEMA_VERSION = 1
BACKUP_ID_RE = re.compile(r"^\d{8}_\d{6}$")
TXT_SECTIONS = ("MANIFEST", "RULES", "GLOSSARY", "EXCEPTIONS", "FEEDBACK", "PRESETS", "STATISTICS")


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def timestamp_id() -> str:
    return datetime.now().strftime("%Y%m%d_%H%M%S")


def ensure_memory_dirs(memory_dir: Path) -> Path:
    backups = memory_dir / "backups"
    backups.mkdir(parents=True, exist_ok=True)
    return backups


def safe_backup_path(memory_dir: Path, backup_id: str) -> Path:
    if not BACKUP_ID_RE.match(backup_id):
        raise ValueError("Invalid backup id.")
    root = ensure_memory_dirs(memory_dir).resolve()
    path = (root / backup_id).resolve()
    if root not in path.parents:
        raise ValueError("Invalid backup path.")
    return path


def _rows(conn, sql: str, params: tuple[Any, ...] = ()) -> list[dict[str, Any]]:
    return [dict(row) for row in conn.execute(sql, params).fetchall()]


def _json_loads(value: str) -> Any:
    try:
        return json.loads(value or "")
    except json.JSONDecodeError:
        return value


def collect_memory(conn, *, scope: str = "all", project: str = "") -> dict[str, Any]:
    project_filter = scope == "current_project" and project
    params = (project,) if project_filter else ()
    project_where = "WHERE project_name = ?" if project_filter else ""
    glossary_where = "WHERE project_id = ?" if project_filter else ""

    rules = _rows(
        conn,
        """
        SELECT rule_id, title AS rule_name, engine, source AS standard, category,
               severity, enabled, config_path AS source, created_at, updated_at
        FROM rule_registry
        ORDER BY rule_id
        """,
    )
    for row in rules:
        row["rule_key"] = f"rule_registry:{row.get('standard') or row.get('engine') or 'registry'}:{row.get('rule_id')}"
        row["promoted_from_issue"] = "Yes"

    glossary = _rows(
        conn,
        f"""
        SELECT id AS term_id, project_id AS project, scope, term_type,
               source_term, preferred_term, description, case_sensitive,
               active, created_at, updated_at
        FROM glossary_terms
        {glossary_where}
        ORDER BY project_id, scope, source_term
        """,
        params,
    )

    feedback = _rows(
        conn,
        f"""
        SELECT rf.rule_id, rf.engine, rf.decision, rf.reviewer_note,
               rf.document_id AS document_reference, rf.created_at
        FROM rule_feedback rf
        LEFT JOIN documents d ON d.id = rf.document_id
        {"WHERE d.project_name = ?" if project_filter else ""}
        ORDER BY rf.created_at DESC
        """,
        params,
    )

    presets = _rows(
        conn,
        f"""
        SELECT project_name AS project, profile_id AS profile,
               enabled_standards_json AS enabled_standards, updated_at
        FROM project_settings
        {project_where}
        ORDER BY project_name
        """,
        params,
    )
    for row in presets:
        row["enabled_engines"] = "Stored in app_metadata review_engine_config if configured"
        row["severity_threshold"] = "Not available"
        row["review_mode"] = "Not available"
        row["enabled_standards"] = json.dumps(_json_loads(row.get("enabled_standards") or "[]"), ensure_ascii=False)

    statistics = _rows(
        conn,
        f"""
        SELECT i.rule_id, i.engine, i.standard, COUNT(*) AS rule_trigger_count,
               SUM(CASE WHEN i.status = 'accepted' THEN 1 ELSE 0 END) AS accepted_count,
               SUM(CASE WHEN i.status = 'rejected' THEN 1 ELSE 0 END) AS rejected_count,
               SUM(CASE WHEN i.status = 'ignored_by_rule' THEN 1 ELSE 0 END) AS ignored_count,
               MAX(i.created_at) AS last_triggered
        FROM issues i
        LEFT JOIN documents d ON d.id = i.document_id
        {"WHERE d.project_name = ?" if project_filter else ""}
        GROUP BY i.rule_id, i.engine, i.standard
        ORDER BY rule_trigger_count DESC
        """,
        params,
    )

    manifest = {
        "schema_version": SCHEMA_VERSION,
        "application_version": "pdf_english_reviewer_260713",
        "created_at": utc_now(),
        "project_scope": project if project_filter else "All Projects",
        "included_sections": "rules,glossary,exceptions,feedback,presets,statistics",
        "record_counts": json.dumps({
            "rules": len(rules),
            "glossary": len(glossary),
            "exceptions": 0,
            "feedback": len(feedback),
            "presets": len(presets),
            "statistics": len(statistics),
        }, ensure_ascii=False),
        "checksum": "",
    }
    memory = {
        "manifest": manifest,
        "rules": rules,
        "glossary": glossary,
        "exceptions": [],
        "feedback": feedback,
        "presets": presets,
        "statistics": statistics,
    }
    manifest["checksum"] = memory_checksum(memory)
    return memory


def memory_checksum(memory: dict[str, Any]) -> str:
    payload = {key: value for key, value in memory.items() if key != "manifest"}
    data = json.dumps(payload, sort_keys=True, ensure_ascii=False, default=str).encode("utf-8")
    return hashlib.sha256(data).hexdigest()


def _escape_txt(value: Any) -> str:
    text = str(safe_excel_value(value))
    return text.replace("\\", "\\\\").replace("\t", "\\t").replace("\r", "\\r").replace("\n", "\\n")


def _unescape_txt(value: str) -> str:
    out = []
    escape = False
    for char in value:
        if escape:
            out.append({"t": "\t", "n": "\n", "r": "\r", "\\": "\\"}.get(char, char))
            escape = False
        elif char == "\\":
            escape = True
        else:
            out.append(char)
    if escape:
        out.append("\\")
    return "".join(out)


def write_memory_txt(path: Path, memory: dict[str, Any]) -> Path:
    lines: list[str] = [
        "# PDF English Reviewer Review Memory",
        "# UTF-8, section headers, first row column headers, tab-delimited, escaped \\t \\n \\r \\\\",
        "",
    ]
    section_map = {
        "MANIFEST": [memory.get("manifest", {})],
        "RULES": memory.get("rules", []),
        "GLOSSARY": memory.get("glossary", []),
        "EXCEPTIONS": memory.get("exceptions", []),
        "FEEDBACK": memory.get("feedback", []),
        "PRESETS": memory.get("presets", []),
        "STATISTICS": memory.get("statistics", []),
    }
    for section, rows in section_map.items():
        lines.append(f"[{section}]")
        headers = sorted({key for row in rows for key in row.keys()}) if rows else ["N/A"]
        lines.append("\t".join(headers))
        for row in rows:
            lines.append("\t".join(_escape_txt(row.get(header, "")) for header in headers))
        lines.append("")
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def create_backup(
    conn,
    *,
    memory_dir: Path,
    scope: str = "all",
    project: str = "",
    team_rules_path: Path | None = None,
    exceptions_path: Path | None = None,
    profiles_path: Path | None = None,
) -> dict[str, Any]:
    backup_id = timestamp_id()
    backup_dir = safe_backup_path(memory_dir, backup_id)
    backup_dir.mkdir(parents=True, exist_ok=False)
    source_dir = backup_dir / "source_files"
    source_dir.mkdir(parents=True, exist_ok=True)
    memory = collect_memory(conn, scope=scope, project=project)
    manifest_path = backup_dir / "manifest.json"
    manifest_path.write_text(json.dumps(memory["manifest"], ensure_ascii=False, indent=2), encoding="utf-8")
    xlsx_path = build_memory_workbook(backup_dir / "review_memory.xlsx", memory)
    txt_path = write_memory_txt(backup_dir / "review_memory.txt", memory)
    for source in (team_rules_path, exceptions_path, profiles_path):
        if source and source.exists() and source.is_file():
            shutil.copy2(source, source_dir / source.name)
    return backup_record(backup_dir)


def backup_record(path: Path) -> dict[str, Any]:
    manifest_path = path / "manifest.json"
    manifest = {}
    if manifest_path.exists():
        try:
            manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        except json.JSONDecodeError:
            manifest = {}
    size = sum(item.stat().st_size for item in path.rglob("*") if item.is_file())
    return {
        "backup_id": path.name,
        "backup_name": path.name,
        "created_at": manifest.get("created_at") or path.name,
        "project_scope": manifest.get("project_scope") or "Not available",
        "file_type": "xlsx, txt, json",
        "record_count": manifest.get("record_counts") or "{}",
        "schema_version": manifest.get("schema_version") or "Not available",
        "file_size": size,
        "status": "Valid" if manifest else "Missing manifest",
        "folder_path": str(path),
    }


def list_backups(memory_dir: Path) -> list[dict[str, Any]]:
    root = ensure_memory_dirs(memory_dir)
    return [backup_record(path) for path in sorted(root.iterdir(), reverse=True) if path.is_dir() and BACKUP_ID_RE.match(path.name)]


def memory_summary(conn, *, memory_dir: Path) -> dict[str, Any]:
    backups = list_backups(memory_dir)
    counts = {
        "total_stored_rules": conn.execute("SELECT COUNT(*) FROM rule_registry").fetchone()[0],
        "promoted_rules": conn.execute("SELECT COUNT(*) FROM issues WHERE promoted_to_rule = 1").fetchone()[0],
        "active_glossary_terms": conn.execute("SELECT COUNT(*) FROM glossary_terms WHERE active = 1").fetchone()[0],
        "rule_exceptions": 0,
        "feedback_records": conn.execute("SELECT COUNT(*) FROM rule_feedback").fetchone()[0],
        "review_presets": conn.execute("SELECT COUNT(*) FROM project_settings").fetchone()[0],
    }
    return {
        **counts,
        "last_backup": backups[0]["created_at"] if backups else "",
        "backup_folder": str(ensure_memory_dirs(memory_dir)),
        "backup_count": len(backups),
    }


def delete_backup(memory_dir: Path, backup_id: str) -> dict[str, Any]:
    path = safe_backup_path(memory_dir, backup_id)
    if not path.exists():
        raise FileNotFoundError("Backup not found.")
    shutil.rmtree(path)
    return {"deleted": backup_id}


def parse_memory_txt(content: str) -> dict[str, list[dict[str, str]]]:
    sections: dict[str, list[dict[str, str]]] = {name: [] for name in TXT_SECTIONS}
    current = ""
    headers: list[str] = []
    for raw_line in content.splitlines():
        line = raw_line.strip("\ufeff")
        if not line or line.startswith("#"):
            continue
        if line.startswith("[") and line.endswith("]"):
            current = line[1:-1].strip().upper()
            headers = []
            if current not in sections:
                sections[current] = []
            continue
        if not current:
            continue
        parts = [_unescape_txt(part) for part in line.split("\t")]
        if not headers:
            headers = parts
            continue
        sections[current].append({headers[index]: parts[index] if index < len(parts) else "" for index in range(len(headers))})
    return sections


def validate_import_file(filename: str, content: bytes) -> dict[str, Any]:
    suffix = Path(filename).suffix.lower()
    if suffix not in {".txt", ".xlsx"}:
        return {"valid": False, "errors": ["Only .txt and .xlsx files are allowed."], "format": suffix}
    if len(content) > 25 * 1024 * 1024:
        return {"valid": False, "errors": ["Import file is larger than 25 MB."], "format": suffix}
    errors: list[str] = []
    sections: dict[str, list[dict[str, str]]] = {}
    if suffix == ".txt":
        sections = parse_memory_txt(content.decode("utf-8-sig"))
        manifest_rows = sections.get("MANIFEST") or []
        schema = ""
        if manifest_rows:
            row = manifest_rows[0]
            schema = str(row.get("schema_version") or row.get("Schema Version") or "")
        if schema and str(schema) != str(SCHEMA_VERSION):
            errors.append(f"Unsupported schema_version: {schema}")
    else:
        try:
            from openpyxl import load_workbook
            workbook = load_workbook(filename=__import__("io").BytesIO(content), read_only=True, data_only=True)
            missing = {"Manifest", "Rules", "Glossary"} - set(workbook.sheetnames)
            if missing:
                errors.append(f"Missing sheets: {', '.join(sorted(missing))}")
        except Exception as exc:  # pragma: no cover - depends on workbook parser
            errors.append(f"XLSX validation failed: {type(exc).__name__}")
    return {
        "valid": not errors,
        "errors": errors,
        "format": suffix.lstrip("."),
        "sections": {key: len(value) for key, value in sections.items()} if sections else {},
    }


def preview_import(conn, *, filename: str, content: bytes) -> dict[str, Any]:
    validation = validate_import_file(filename, content)
    if not validation["valid"]:
        return {"valid": False, "validation": validation, "new_records": 0, "updated_records": 0, "conflicts": 0, "invalid_records": 0}
    suffix = Path(filename).suffix.lower()
    sections = parse_memory_txt(content.decode("utf-8-sig")) if suffix == ".txt" else {}
    glossary = sections.get("GLOSSARY", [])
    rules = sections.get("RULES", [])
    existing_terms = {
        row[0] for row in conn.execute("SELECT id FROM glossary_terms").fetchall()
    }
    existing_rules = {
        row[0] for row in conn.execute("SELECT rule_id FROM rule_registry").fetchall()
    }
    conflicts = sum(1 for item in glossary if item.get("term_id") in existing_terms)
    conflicts += sum(1 for item in rules if item.get("rule_id") in existing_rules)
    return {
        "valid": True,
        "validation": validation,
        "new_records": max(0, len(glossary) + len(rules) - conflicts),
        "updated_records": 0,
        "unchanged_records": 0,
        "invalid_records": 0,
        "conflicts": conflicts,
        "disabled_records": sum(1 for item in glossary if str(item.get("active", "1")).lower() in {"0", "false", "no"}),
        "missing_references": 0,
        "sections": validation.get("sections", {}),
        "default_conflict_policy": "Skip Existing",
    }


def apply_import(conn, *, filename: str, content: bytes, conflict_policy: str = "skip") -> dict[str, Any]:
    suffix = Path(filename).suffix.lower()
    if suffix != ".txt":
        return {"applied": False, "detail": "XLSX import preview is supported; apply currently requires TXT format."}
    preview = preview_import(conn, filename=filename, content=content)
    if not preview.get("valid"):
        return {"applied": False, "preview": preview}
    sections = parse_memory_txt(content.decode("utf-8-sig"))
    now = utc_now()
    inserted = 0
    skipped = 0
    with conn:
        for row in sections.get("GLOSSARY", []):
            term_id = row.get("term_id") or row.get("id")
            if not term_id:
                skipped += 1
                continue
            exists = conn.execute("SELECT 1 FROM glossary_terms WHERE id = ?", (term_id,)).fetchone()
            if exists and conflict_policy == "skip":
                skipped += 1
                continue
            conn.execute(
                """
                INSERT INTO glossary_terms (
                    id, project_id, scope, term_type, source_term, preferred_term,
                    description, case_sensitive, active, created_at, updated_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                ON CONFLICT(id) DO UPDATE SET
                    preferred_term = excluded.preferred_term,
                    description = excluded.description,
                    case_sensitive = excluded.case_sensitive,
                    active = excluded.active,
                    updated_at = excluded.updated_at
                """,
                (
                    term_id,
                    row.get("project") or "Common Glossary",
                    row.get("scope") or "project",
                    row.get("term_type") or "protected",
                    row.get("source_term") or "",
                    row.get("preferred_term") or "",
                    row.get("description") or "",
                    int(str(row.get("case_sensitive") or "0").lower() in {"1", "true", "yes"}),
                    int(str(row.get("active") or "1").lower() in {"1", "true", "yes"}),
                    row.get("created_at") or now,
                    now,
                ),
            )
            inserted += 1
        for row in sections.get("RULES", []):
            rule_id = row.get("rule_id") or ""
            if not rule_id:
                skipped += 1
                continue
            exists = conn.execute("SELECT 1 FROM rule_registry WHERE rule_id = ?", (rule_id,)).fetchone()
            if exists and conflict_policy == "skip":
                skipped += 1
                continue
            conn.execute(
                """
                INSERT INTO rule_registry (rule_id, title, engine, category, severity, source, enabled, config_path, created_at, updated_at)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                ON CONFLICT(rule_id) DO UPDATE SET
                    title = excluded.title,
                    category = excluded.category,
                    severity = excluded.severity,
                    enabled = excluded.enabled,
                    updated_at = excluded.updated_at
                """,
                (
                    rule_id,
                    row.get("rule_name") or rule_id,
                    row.get("engine") or "rule_registry",
                    row.get("category") or "",
                    row.get("severity") or "minor",
                    row.get("standard") or row.get("source") or "ReviewMemory",
                    int(str(row.get("enabled") or "1").lower() in {"1", "true", "yes"}),
                    "review_memory_import",
                    row.get("created_at") or now,
                    now,
                ),
            )
            inserted += 1
    return {"applied": True, "inserted_or_updated": inserted, "skipped": skipped, "conflict_policy": conflict_policy}
