from __future__ import annotations

from copy import copy
from datetime import datetime, timezone
from pathlib import Path
from typing import Any


FORMULA_PREFIXES = ("=", "+", "-", "@")


def safe_excel_value(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, (int, float, bool)):
        return value
    text = str(value)
    if text.startswith(FORMULA_PREFIXES):
        return "'" + text
    return text


def timestamp_name(prefix: str, suffix: str) -> str:
    return f"{prefix}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.{suffix}"


def _append_rows(sheet, headers: list[str], rows: list[dict[str, Any]]) -> None:
    sheet.append(headers)
    for row in rows:
        sheet.append([safe_excel_value(row.get(header, "")) for header in headers])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        font = copy(cell.font)
        font.bold = True
        cell.font = font
    for column_cells in sheet.columns:
        header = str(column_cells[0].value or "")
        width = min(48, max(12, len(header) + 2))
        sheet.column_dimensions[column_cells[0].column_letter].width = width
        for cell in column_cells:
            alignment = copy(cell.alignment)
            alignment.wrap_text = True
            alignment.vertical = "top"
            cell.alignment = alignment


def build_rule_dashboard_workbook(
    path: Path,
    *,
    summary: dict[str, Any],
    inventory: list[dict[str, Any]],
    findings: list[dict[str, Any]],
    statistics: list[dict[str, Any]],
    criteria: dict[str, Any],
) -> Path:
    try:
        from openpyxl import Workbook
    except ImportError as exc:  # pragma: no cover - depends on local environment
        raise RuntimeError("openpyxl is required for .xlsx export.") from exc

    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Summary"
    summary_rows = [
        {"Metric": "Export Date", "Value": datetime.now(timezone.utc).isoformat()},
        *[{"Metric": key, "Value": value} for key, value in summary.items() if not isinstance(value, (dict, list))],
    ]
    _append_rows(summary_sheet, ["Metric", "Value"], summary_rows)

    inventory_headers = [
        "rule_key", "rule_id", "rule_name", "rule_source", "standard", "engine",
        "category", "severity", "description", "pattern_type", "pattern",
        "replacement", "enabled", "configuration_source", "trigger_count",
        "affected_documents", "affected_projects", "open_count", "accepted_count",
        "rejected_count", "ignored_count", "first_triggered", "last_triggered",
        "rule_status",
    ]
    _append_rows(workbook.create_sheet("Rule Inventory"), inventory_headers, inventory)

    finding_headers = [
        "document", "project", "review_session", "review_date", "page",
        "rule_key", "rule_id", "rule_name", "rule_source", "engine", "standard",
        "category", "severity", "source_text", "suggested_text", "message",
        "status", "reviewer", "reviewer_comment", "detected_at",
    ]
    _append_rows(workbook.create_sheet("Rule Findings"), finding_headers, findings)

    stat_headers = [
        "rule_key", "rule_id", "rule_name", "rule_source", "engine", "standard",
        "trigger_count", "affected_documents", "affected_projects", "open_count",
        "accepted_count", "rejected_count", "ignored_count", "acceptance_rate",
        "rejection_rate", "ignore_rate", "first_triggered", "last_triggered",
    ]
    stat_rows = []
    inventory_by_key = {item["rule_key"]: item for item in inventory}
    for stat in statistics:
        rule = inventory_by_key.get(stat["rule_key"], {})
        row = {**stat}
        row.update({
            "rule_id": rule.get("rule_id", ""),
            "rule_name": rule.get("rule_name", ""),
            "rule_source": rule.get("rule_source", ""),
            "engine": rule.get("engine", ""),
            "standard": rule.get("standard", ""),
        })
        stat_rows.append(row)
    _append_rows(workbook.create_sheet("Rule Statistics"), stat_headers, stat_rows)

    criteria_rows = [{"Criterion": key, "Value": value} for key, value in criteria.items()]
    criteria_rows.append({"Criterion": "Export Date", "Value": datetime.now(timezone.utc).isoformat()})
    _append_rows(workbook.create_sheet("Export Criteria"), ["Criterion", "Value"], criteria_rows)

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


def build_memory_workbook(path: Path, memory: dict[str, Any]) -> Path:
    try:
        from openpyxl import Workbook
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("openpyxl is required for .xlsx export.") from exc

    workbook = Workbook()
    manifest = workbook.active
    manifest.title = "Manifest"
    _append_rows(manifest, ["Key", "Value"], [{"Key": k, "Value": v} for k, v in memory["manifest"].items()])
    sheets = {
        "Rules": memory.get("rules", []),
        "Glossary": memory.get("glossary", []),
        "Rule Exceptions": memory.get("exceptions", []),
        "Feedback": memory.get("feedback", []),
        "Review Presets": memory.get("presets", []),
        "Statistics": memory.get("statistics", []),
    }
    for title, rows in sheets.items():
        headers = sorted({key for row in rows for key in row.keys()}) if rows else ["N/A"]
        _append_rows(workbook.create_sheet(title), headers, rows)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path
